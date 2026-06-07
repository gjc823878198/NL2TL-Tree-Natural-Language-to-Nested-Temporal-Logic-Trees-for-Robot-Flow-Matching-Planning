"""Build Figure 3 as ONE editable PPTX (figure3.pptx): panels (a) and (b) as
native PowerPoint shapes, panel (c) as the embedded trajectory plot + re-typeset
text, all in ONE consistent font/size.  Export -> figure3.pdf for the paper.

    python3 build_figure3.py        # writes figure3.pptx  (then soffice -> pdf)
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ASSET = HERE / "fig3_assets"
FONT = "Arial"
# palette (matches the polished matplotlib figure)
def C(h): return RGBColor.from_string(h)
INK=C("16213A"); OPN=C("D7E3F2"); OPE=C("2C3E6B"); GRY=C("E3E6EA"); GRE=C("7A7A7A")
REACH=C("3A6EA5"); AVOID=C("5B6675"); MECH=C("FFF4D6"); MECHE=C("B8860B")
GATEF=C("CFE0D2"); GATEE=C("4F9E74"); REDF=C("FBE2DF"); REDE=C("C97A6A")
GREY9=C("888888"); CAP=C("444444")
# font sizes (ONE consistent scheme across a/b/c)
F_TITLE=11; F_NODE=8; F_LBL=8; F_CAP=7.5

prs = Presentation()
prs.slide_width = Inches(13.33); prs.slide_height = Inches(2.85)
slide = prs.slides.add_slide(prs.slide_layouts[6])
S = slide.shapes

def noshadow(sp): sp.shadow.inherit=False

def txt(x,y,w,h,lines,size=F_LBL,bold=True,color=INK,align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,italic=False):
    tb=S.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True
    for m in ("left","right","top","bottom"): setattr(tf,f"margin_{m}",0)
    tf.vertical_anchor=anchor
    if isinstance(lines,str): lines=lines.split("\n")
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=ln; f=r.font
        f.size=Pt(size); f.bold=bold; f.italic=italic; f.name=FONT; f.color.rgb=color
    return tb

def oval(cx,cy,d,fill,line,lw=1.0,tcolor=INK,label="",fs=F_NODE):
    sp=S.add_shape(MSO_SHAPE.OVAL,Inches(cx-d/2),Inches(cy-d/2),Inches(d),Inches(d))
    sp.fill.solid(); sp.fill.fore_color.rgb=fill; sp.line.color.rgb=line
    sp.line.width=Pt(lw); noshadow(sp)
    if label:
        tf=sp.text_frame; tf.word_wrap=False
        for m in ("left","right","top","bottom"): setattr(tf,f"margin_{m}",0)
        p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
        r=p.add_run(); r.text=label; f=r.font
        f.size=Pt(fs); f.bold=True; f.name=FONT; f.color.rgb=tcolor
    return sp

def rrect(x,y,w,h,fill,line,lw=1.2,lines="",fs=F_LBL,tcolor=INK,bold=True):
    sp=S.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sp.adjustments[0]=0.16
    sp.fill.solid(); sp.fill.fore_color.rgb=fill; sp.line.color.rgb=line
    sp.line.width=Pt(lw); noshadow(sp)
    if lines:
        tf=sp.text_frame; tf.word_wrap=True
        for m in ("left","right","top","bottom"): setattr(tf,f"margin_{m}",Pt(1))
        tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        if isinstance(lines,str): lines=lines.split("\n")
        for i,ln in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=PP_ALIGN.CENTER
            r=p.add_run(); r.text=ln; f=r.font
            f.size=Pt(fs); f.bold=bold; f.name=FONT; f.color.rgb=tcolor
    return sp

def conn(x0,y0,x1,y1,color=GREY9,lw=1.1,arrow=False):
    cn=S.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x0),Inches(y0),Inches(x1),Inches(y1))
    cn.line.color.rgb=color; cn.line.width=Pt(lw); noshadow(cn)
    if arrow:
        ln=cn.line._get_or_add_ln()
        te=ln.makeelement(qn("a:tailEnd"),{"type":"triangle","w":"med","len":"med"})
        ln.append(te)
    return cn

def pic(path,x,y,w=None,h=None):
    return S.add_picture(str(path),Inches(x),Inches(y),
                         Inches(w) if w else None, Inches(h) if h else None)

# ============================ panel (a): STL tree ============================
ax0,ay0,aw,ah = 0.18,0.46,2.55,2.00     # tree drawing area
def A(mx,my): return ax0+(mx/10)*aw, ay0+(1-my/10)*ah
txt(0.10,0.06,2.7,0.32,"(a) Nested STL spec",size=F_TITLE,align=PP_ALIGN.LEFT)
nodes={"and":(5,9,"∧",OPN,INK),"F":(2.4,6.4,"F[a,b]",OPN,INK),
       "g":(2.4,3.6,"reach g",REACH,C("FFFFFF")),"G1":(5,6.4,"G",OPN,INK),
       "n1":(5,4.2,"¬",GRY,INK),"o":(5,2.0,"avoid o",AVOID,C("FFFFFF")),
       "Gks":(7.8,6.4,"G",OPN,INK),"nks":(7.8,4.2,"¬",GRY,INK),
       "u":(7.8,2.0,"unsafe",AVOID,C("FFFFFF"))}
edges=[("and","F"),("F","g"),("and","G1"),("G1","n1"),("n1","o"),
       ("and","Gks"),("Gks","nks"),("nks","u")]
D=0.40
for a,b in edges:
    x0,y0=A(nodes[a][0],nodes[a][1]); x1,y1=A(nodes[b][0],nodes[b][1])
    conn(x0,y0,x1,y1,color=GREY9,lw=1.1)
for k,(mx,my,lb,fc,tc) in nodes.items():
    cx,cy=A(mx,my); oval(cx,cy,D,fc,C("333333"),lw=1.0,tcolor=tc,label=lb)
txt(0.10,2.55,2.7,0.28,"keep-safe = G ¬ unsafe, grounded by sensed disks",
    size=F_CAP,color=CAP,align=PP_ALIGN.CENTER)

# ============================ panel (b): planner =============================
bx0,by0,bw,bh = 3.05,0.0,4.85,2.85
def B(mx,my): return bx0+(mx/100)*bw, by0+(1-my/100)*bh
txt(bx0,0.06,bw,0.32,"(b) Frozen flow-matching planner + robustness guidance",
    size=F_TITLE,align=PP_ALIGN.LEFT)
# encoder box
ex,ey=B(1,84); ex2,ey2=B(28,64); rrect(ex,ey,ex2-ex,ey2-ey,OPN,OPE,lw=1.3,
    lines="frozen GNN\nencoder",fs=F_NODE)
cxz,cyz=B(14.5,58.5); txt(cxz-0.6,cyz-0.12,1.2,0.24,"⇒ conditioning z",size=F_CAP)
# arrow encoder -> flow
a1x,a1y=B(28,74); a2x,a2y=B(41.4,74); conn(a1x,a1y,a2x,a2y,color=OPE,lw=1.4,arrow=True)
# flow row (title already says "flow-matching"; no extra row label to avoid clutter)
xs=[47,62,77,92]; labels=["xT","xt","⋯","x0"]; yr=74; cd=0.46
for i,(mx,lb) in enumerate(zip(xs,labels)):
    cx,cy=B(mx,yr); oval(cx,cy,cd,GATEF,GATEE,lw=1.3,tcolor=INK,label=lb)
    if i<3:
        x0,_=B(xs[i]+6,yr); x1,_=B(xs[i+1]-6,yr); conn(x0,cy,x1,cy,color=GATEE,lw=1.3,arrow=True)
# guidance box
gx,gy=B(33,56); gx2,gy2=B(98,44); rrect(gx,gy,gx2-gx,gy2-gy,REDF,REDE,lw=1.4,
    lines="differentiable STL robustness ∇ρ (STLCG) steers sampling",fs=F_CAP)
for mx in xs[:-1]:
    x,y0=B(mx,56); _,y1=B(mx,yr); cyc=B(mx,yr)[1]; conn(x,gy,x,cyc+cd/2,color=REDE,lw=1.2,arrow=True)
# mechanism tags
tags=[("operator normalization\n→,↔ ⇒ basis",1,34),
      ("nested-tree\ndecomposition",37.5,24),
      ("feasibility self-check\n→ A* if OOD",64,34)]
for t,mx,mw in tags:
    x,y=B(mx,37); x2,y2=B(mx+mw,22); rrect(x,y,x2-x,y2-y,MECH,MECHE,lw=1.2,lines=t,fs=F_CAP)
# formula (image)
fx,fy=B(50,14); pic(ASSET/"b_formula.png",bx0+0.25,fy-0.10,w=bw-0.5)

# ============================ panel (c): trajectory =========================
cx0=8.05
txt(cx0,0.06,5.0,0.32,"(c) TeLoGraF plan (no A*)",size=F_TITLE,align=PP_ALIGN.LEFT)
# embedded plot (keep aspect 1130x867)
pic(ASSET/"c_plot.png",cx0+0.02,0.44,h=2.30)
tx=cx0+3.20
txt(tx,0.55,2.05,0.95,["NL: “Reach the goal within","70 seconds while always","keeping safe from every","obstacle.”"],
    size=F_LBL,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,italic=True)
pic(ASSET/"c_stl.png",tx,1.55,w=1.95)
txt(tx,1.95,2.1,0.24,"exec 53 s + 4 s plan (in time)",size=F_LBL,align=PP_ALIGN.LEFT)
pic(ASSET/"c_rho.png",tx,2.25,w=2.0)

prs.save(str(HERE/"figure3.pptx"))
print("wrote", HERE/"figure3.pptx")
